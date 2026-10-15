"""Extract the three reported city-index tables from the competition report.

The source DOCX is intentionally not distributed. This migration helper is retained
to document how the public processed CSV was produced from Tables 3--5.
"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

from docx import Document


YEARS = [2010, 2015, 2020, 2022]
CITY_NAMES = {
    "郑州市": "Zhengzhou",
    "开封市": "Kaifeng",
    "洛阳市": "Luoyang",
    "平顶山市": "Pingdingshan",
    "安阳市": "Anyang",
    "鹤壁市": "Hebi",
    "新乡市": "Xinxiang",
    "焦作市": "Jiaozuo",
    "濮阳市": "Puyang",
    "许昌市": "Xuchang",
    "漯河市": "Luohe",
    "三门峡市": "Sanmenxia",
    "南阳市": "Nanyang",
    "商丘市": "Shangqiu",
    "信阳市": "Xinyang",
    "周口市": "Zhoukou",
    "驻马店市": "Zhumadian",
}


def normalise_city(value: str) -> str:
    return re.sub(r"\s+", "", value)


def read_table(table) -> dict[str, list[float]]:
    records: dict[str, list[float]] = {}
    for row in table.rows[1:]:
        values = [cell.text.strip() for cell in row.cells]
        city_zh = normalise_city(values[0])
        if not city_zh:
            continue
        if city_zh not in CITY_NAMES:
            raise ValueError(f"Unexpected city label: {values[0]!r}")
        records[city_zh] = [float(value) for value in values[1:5]]
    return records


def coordination_class(value: float) -> str:
    if value <= 0.2:
        return "severe imbalance"
    if value <= 0.4:
        return "moderate imbalance"
    if value <= 0.6:
        return "basic coordination"
    if value <= 0.8:
        return "moderate coordination"
    return "high coordination"


def migrate(source: Path, destination: Path) -> None:
    document = Document(source)
    if len(document.tables) < 5:
        raise ValueError("Expected at least five tables in the source report")

    urbanisation = read_table(document.tables[2])
    ecosystem_health = read_table(document.tables[3])
    coordination = read_table(document.tables[4])
    city_sets = [set(urbanisation), set(ecosystem_health), set(coordination)]
    if not city_sets[0] == city_sets[1] == city_sets[2] == set(CITY_NAMES):
        raise ValueError("The three result tables do not contain the same 17 cities")

    destination.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "year",
        "city",
        "city_zh",
        "urbanisation_index_reported",
        "ecosystem_health_index_reported",
        "coupling_coordination_reported",
        "coordination_class",
    ]
    with destination.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames)
        writer.writeheader()
        for city_zh, city in CITY_NAMES.items():
            for index, year in enumerate(YEARS):
                coupling = coordination[city_zh][index]
                writer.writerow(
                    {
                        "year": year,
                        "city": city,
                        "city_zh": city_zh,
                        "urbanisation_index_reported": urbanisation[city_zh][index],
                        "ecosystem_health_index_reported": ecosystem_health[city_zh][index],
                        "coupling_coordination_reported": coupling,
                        "coordination_class": coordination_class(coupling),
                    }
                )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source_docx", type=Path)
    parser.add_argument("output_csv", type=Path)
    args = parser.parse_args()
    migrate(args.source_docx, args.output_csv)


if __name__ == "__main__":
    main()
